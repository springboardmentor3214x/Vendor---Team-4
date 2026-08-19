from io import BytesIO
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_value(value):

    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")

    return value


def generate_excel(
    report_title: str,
    report_data: dict,
    filters: dict = None
):

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Report"

    # ---------------------------------------
    # Styles
    # ---------------------------------------

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="5B9BD5"
    )

    title_font = Font(
        bold=True,
        color="FFFFFF",
        size=16
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        left=Side(style="thin", color="D9E1F2"),
        right=Side(style="thin", color="D9E1F2"),
        top=Side(style="thin", color="D9E1F2"),
        bottom=Side(style="thin", color="D9E1F2")
    )

    # ---------------------------------------
    # Title
    # ---------------------------------------

    worksheet.merge_cells(
        "A1:F1"
    )

    worksheet["A1"] = report_title
    worksheet["A1"].font = title_font
    worksheet["A1"].fill = title_fill
    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    # ---------------------------------------
    # Generation date
    # ---------------------------------------

    worksheet["A3"] = "Generated On"
    worksheet["B3"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M"
    )

    worksheet["A3"].font = Font(bold=True)

    # ---------------------------------------
    # Filters
    # ---------------------------------------

    current_row = 5

    if filters:

        active_filters = {
            key: value
            for key, value in filters.items()
            if value is not None and value != ""
        }

        if active_filters:

            worksheet.cell(
                row=current_row,
                column=1,
                value="Applied Filters"
            )

            worksheet.cell(
                row=current_row,
                column=1
            ).font = Font(
                bold=True,
                size=12
            )

            current_row += 1

            worksheet.cell(
                row=current_row,
                column=1,
                value="Filter"
            )

            worksheet.cell(
                row=current_row,
                column=2,
                value="Value"
            )

            for cell in worksheet[current_row]:

                if cell.column <= 2:
                    cell.fill = header_fill
                    cell.font = header_font

            current_row += 1

            for key, value in active_filters.items():

                worksheet.cell(
                    row=current_row,
                    column=1,
                    value=key.replace(
                        "_", " "
                    ).title()
                )

                worksheet.cell(
                    row=current_row,
                    column=2,
                    value=format_value(value)
                )

                current_row += 1

            current_row += 2

    # ---------------------------------------
    # Report data
    # ---------------------------------------

    for section_name, section_data in report_data.items():

        if section_name == "report_type":
            continue

        worksheet.cell(
            row=current_row,
            column=1,
            value=section_name.replace(
                "_", " "
            ).title()
        )

        worksheet.cell(
            row=current_row,
            column=1
        ).font = Font(
            bold=True,
            size=12
        )

        current_row += 1

        # -----------------------------------
        # Dictionary
        # -----------------------------------

        if isinstance(section_data, dict):

            worksheet.cell(
                row=current_row,
                column=1,
                value="Metric"
            )

            worksheet.cell(
                row=current_row,
                column=2,
                value="Value"
            )

            for col in range(1, 3):

                cell = worksheet.cell(
                    row=current_row,
                    column=col
                )

                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border

            current_row += 1

            for key, value in section_data.items():

                worksheet.cell(
                    row=current_row,
                    column=1,
                    value=key.replace(
                        "_", " "
                    ).title()
                )

                if isinstance(value, dict):

                    worksheet.cell(
                        row=current_row,
                        column=2,
                        value=str(value)
                    )

                elif isinstance(value, list):

                    worksheet.cell(
                        row=current_row,
                        column=2,
                        value=str(value)
                    )

                else:

                    worksheet.cell(
                        row=current_row,
                        column=2,
                        value=format_value(value)
                    )

                for col in range(1, 3):

                    worksheet.cell(
                        row=current_row,
                        column=col
                    ).border = thin_border

                current_row += 1

        # -----------------------------------
        # List of records
        # -----------------------------------

        elif isinstance(section_data, list):

            if section_data:

                first_item = section_data[0]

                if isinstance(first_item, dict):

                    keys = list(
                        first_item.keys()
                    )

                    # Header
                    for column_index, key in enumerate(
                        keys,
                        start=1
                    ):

                        cell = worksheet.cell(
                            row=current_row,
                            column=column_index,
                            value=key.replace(
                                "_", " "
                            ).title()
                        )

                        cell.fill = header_fill
                        cell.font = header_font
                        cell.border = thin_border
                        cell.alignment = Alignment(
                            horizontal="center"
                        )

                    current_row += 1

                    # Data
                    for item in section_data:

                        for column_index, key in enumerate(
                            keys,
                            start=1
                        ):

                            cell = worksheet.cell(
                                row=current_row,
                                column=column_index,
                                value=format_value(
                                    item.get(key)
                                )
                            )

                            cell.border = thin_border
                            cell.alignment = Alignment(
                                vertical="top",
                                wrap_text=True
                            )

                        current_row += 1

        else:

            worksheet.cell(
                row=current_row,
                column=1,
                value=format_value(
                    section_data
                )
            )

            current_row += 1

        current_row += 2

    # ---------------------------------------
    # Formatting
    # ---------------------------------------

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:
                cell_length = len(
                    str(cell.value)
                )

                if cell_length > max_length:
                    max_length = cell_length

            except Exception:
                pass

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            40
        )

    worksheet.freeze_panes = "A6"

    # ---------------------------------------
    # Save to memory
    # ---------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output
